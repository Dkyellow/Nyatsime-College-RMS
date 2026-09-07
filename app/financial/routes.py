import csv
import io
from datetime import date
from flask import render_template, redirect, url_for, flash, request
from flask_login import login_required, current_user
from functools import wraps
from app.financial import financial_bp
from app.models import db, Student, Payment, Class, AcademicYear, AuditLog
from app.services import periods


def admin_required(f):
    @wraps(f)
    @login_required
    def decorated_function(*args, **kwargs):
        if current_user.role != 'admin':
            flash('Access denied.', 'danger')
            return redirect(url_for('auth.login'))
        return f(*args, **kwargs)
    return decorated_function


def log_action(action, resource_type, resource_id=None, details=''):
    entry = AuditLog(user_id=current_user.id, action=action,
                     resource_type=resource_type, resource_id=resource_id,
                     details=details, ip_address=request.remote_addr)
    db.session.add(entry)


@financial_bp.route('/')
@admin_required
def index():
    default_term, default_year = periods.get_default_period()
    academic_term = request.args.get('term', default_term)
    academic_year = request.args.get('year', default_year)
    class_id = request.args.get('class_id', type=int)
    status_filter = request.args.get('status', '')

    query = Student.query.filter_by(is_active=True)
    if class_id:
        query = query.filter_by(class_id=class_id)

    students = query.order_by(Student.admission_number).all()

    payments = {}
    p_query = Payment.query.filter_by(academic_year=academic_year, academic_term=academic_term)
    for p in p_query.all():
        payments[p.student_id] = p

    student_data = []
    for s in students:
        p = payments.get(s.id)
        paid = p.is_paid if p else False
        if status_filter == 'paid' and not paid:
            continue
        if status_filter == 'unpaid' and paid:
            continue
        student_data.append({
            'student': s,
            'payment': p,
            'is_paid': paid,
        })

    classes = Class.query.order_by(Class.name).all()
    years = sorted({y[0] for y in db.session.query(Payment.academic_year).distinct()
                    if y[0]}, reverse=True)
    if academic_year and academic_year not in years:
        years.insert(0, academic_year)

    total_paid = sum(1 for d in student_data if d['is_paid'])
    total_unpaid = len(student_data) - total_paid

    return render_template('admin/financial.html',
                           student_data=student_data, classes=classes,
                           years=years, academic_term=academic_term,
                           academic_year=academic_year,
                           selected_class=class_id, status_filter=status_filter,
                           total_paid=total_paid, total_unpaid=total_unpaid)


@financial_bp.route('/upload', methods=['POST'])
@admin_required
def upload():
    default_term, default_year = periods.get_default_period()
    academic_term = request.form.get('term', default_term)
    academic_year = request.form.get('year', default_year)

    if 'file' not in request.files or request.files['file'].filename == '':
        flash('No file selected.', 'danger')
        return redirect(url_for('financial.index', term=academic_term, year=academic_year))

    file = request.files['file']
    try:
        if file.filename.endswith('.csv'):
            content = file.read().decode('utf-8')
            reader = csv.DictReader(io.StringIO(content))
        elif file.filename.endswith(('.xlsx', '.xls')):
            import openpyxl
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            headers = [cell.value for cell in ws[1]]
            reader = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                reader.append(dict(zip(headers, row)))
        else:
            flash('Unsupported file format. Please use CSV or Excel.', 'danger')
            return redirect(url_for('financial.index', term=academic_term, year=academic_year))

        count = 0
        for row in reader:
            admission_number = str(row.get('admission_number', '')).strip()
            if not admission_number:
                continue

            student = Student.query.filter_by(admission_number=admission_number).first()
            if not student:
                continue

            payment = Payment.query.filter_by(
                student_id=student.id, academic_year=academic_year,
                academic_term=academic_term).first()
            if not payment:
                payment = Payment(student_id=student.id, academic_year=academic_year,
                                  academic_term=academic_term)
                db.session.add(payment)

            amount_paid_raw = row.get('amount_paid', row.get('amount', ''))
            amount_due_raw = row.get('amount_due', row.get('fee', ''))
            receipt_raw = str(row.get('receipt_number', row.get('receipt', ''))).strip()
            is_paid_raw = str(row.get('is_paid', row.get('paid', ''))).strip().lower()

            if amount_paid_raw not in (None, ''):
                try:
                    payment.amount_paid = float(amount_paid_raw)
                except (ValueError, TypeError):
                    pass
            if amount_due_raw not in (None, ''):
                try:
                    payment.amount_due = float(amount_due_raw)
                except (ValueError, TypeError):
                    pass
            if receipt_raw:
                payment.receipt_number = receipt_raw

            if is_paid_raw in ('1', 'true', 'yes', 'paid'):
                payment.is_paid = True
            elif is_paid_raw in ('0', 'false', 'no', 'unpaid'):
                payment.is_paid = False
            elif payment.amount_paid >= payment.amount_due and payment.amount_due > 0:
                payment.is_paid = True

            if not payment.payment_date:
                payment.payment_date = date.today()

            count += 1

        db.session.commit()
        log_action('upload', 'payment', None, f'Uploaded {count} payment records for {academic_term} {academic_year}')
        flash(f'{count} payment records uploaded successfully.', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Error uploading payments: {str(e)}', 'danger')

    return redirect(url_for('financial.index', term=academic_term, year=academic_year))


@financial_bp.route('/toggle/<int:student_id>', methods=['POST'])
@admin_required
def toggle_payment(student_id):
    default_term, default_year = periods.get_default_period()
    academic_term = request.form.get('term', default_term)
    academic_year = request.form.get('year', default_year)

    student = Student.query.get_or_404(student_id)
    payment = Payment.query.filter_by(
        student_id=student_id, academic_year=academic_year,
        academic_term=academic_term).first()

    if not payment:
        payment = Payment(student_id=student_id, academic_year=academic_year,
                          academic_term=academic_term, is_paid=True,
                          payment_date=date.today())
        db.session.add(payment)
    else:
        payment.is_paid = not payment.is_paid
        if payment.is_paid and not payment.payment_date:
            payment.payment_date = date.today()

    db.session.commit()
    status = 'paid' if payment.is_paid else 'unpaid'
    log_action('toggle', 'payment', student_id,
               f'Marked {student.first_name} {student.last_name} as {status}')
    flash(f'{student.first_name} {student.last_name} marked as {status}.', 'success')

    return redirect(url_for('financial.index', term=academic_term, year=academic_year))


@financial_bp.route('/bulk-mark', methods=['POST'])
@admin_required
def bulk_mark():
    default_term, default_year = periods.get_default_period()
    academic_term = request.form.get('term', default_term)
    academic_year = request.form.get('year', default_year)
    action = request.form.get('bulk_action', '')
    selected_ids = request.form.getlist('selected_students')

    if not selected_ids:
        flash('No students selected.', 'warning')
        return redirect(url_for('financial.index', term=academic_term, year=academic_year))

    count = 0
    for sid in selected_ids:
        try:
            student_id = int(sid)
        except (ValueError, TypeError):
            continue

        payment = Payment.query.filter_by(
            student_id=student_id, academic_year=academic_year,
            academic_term=academic_term).first()
        if not payment:
            payment = Payment(student_id=student_id, academic_year=academic_year,
                              academic_term=academic_term)
            db.session.add(payment)

        if action == 'mark_paid':
            payment.is_paid = True
            if not payment.payment_date:
                payment.payment_date = date.today()
        elif action == 'mark_unpaid':
            payment.is_paid = False

        count += 1

    db.session.commit()
    verb = 'paid' if action == 'mark_paid' else 'unpaid'
    log_action('bulk_mark', 'payment', None, f'Bulk marked {count} students as {verb}')
    flash(f'{count} students marked as {verb}.', 'success')

    return redirect(url_for('financial.index', term=academic_term, year=academic_year))
